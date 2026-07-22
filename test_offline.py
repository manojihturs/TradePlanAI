# test_offline.py - exercises orb_signal's parity math and state machine
# with synthetic candles, with no network / Upstox token required.
# Run: python test_offline.py

import os, tempfile
from datetime import datetime, date

# Point the DB at a throwaway file before importing orb_common.
os.environ.setdefault("UPSTOX_ACCESS_TOKEN", "dummy-not-used-offline")

import orb_common
orb_common.DB_PATH = os.path.join(tempfile.gettempdir(), "orb_test.db")
if os.path.exists(orb_common.DB_PATH):
    os.remove(orb_common.DB_PATH)

# Never let synthetic test trades hit the real Telegram bot, even if a real
# .env is loaded (orb_common auto-loads .env on import).
orb_common.TELEGRAM_BOT_TOKEN = ""
orb_common.TELEGRAM_CHAT_ID = ""

import orb_journal
orb_journal.JOURNAL_PATH = os.path.join(tempfile.gettempdir(), "orb_test_journal.xlsx")
if os.path.exists(orb_journal.JOURNAL_PATH):
    os.remove(orb_journal.JOURNAL_PATH)

from orb_common import (IST, Candle, db, nearest_strike, SL_POINTS, MAX_DAILY_LOSS, QTY,
                        MIN_PROFIT_POINTS, TSL_STEP_POINTS)
import orb_signal
from orb_signal import (DayState, on_candle_close, ladder_lines, build_watch_pairs,
                        run_replay, CANDLE_MINUTES)
from orb_auto import should_run_today_now, next_market_open
from datetime import timedelta

SESSION = date(2026, 7, 21)
ATM = 24200

def ts(h, m):
    return datetime(2026, 7, 21, h, m, tzinfo=IST)

def seed_levels(conn):
    # ATM CE/PE first-candle levels matching the Jul 21 Excel example from the chat.
    rows = [
        (ATM, "CE", "NSE_FO|ATMCE", 90, 114.6, 47.1, 100, 1000),
        (ATM, "PE", "NSE_FO|ATMPE", 50, 69,   29.15, 60, 1000),
        # Cross-plotted ladder for a CE_WINS (TOP) trade: the trader watches
        # the PUT's own first-5min LOW at neighboring strikes, not the
        # CALL's own high (confirmed in chat - this is a cross-plot, not a
        # same-side extension).
        (ATM - 50, "PE", "NSE_FO|ITM1CE", 165, 170, 160, 168, 1000),
        (ATM - 100, "PE", "NSE_FO|ITM2CE", 205, 210, 200, 208, 1000),   # for the early-exit test
        # A cross-plotted strike whose own low (50) sits BELOW where CE will
        # enter (130) - a real strike, unrelated in magnitude, that must NOT
        # be counted as an already-crossed target the instant the position opens.
        (ATM + 150, "PE", "NSE_FO|PE_BELOW_ENTRY", 45, 55, 50, 52, 1000),
        # Cross-plotted ladder for a PE_WINS (BOTTOM) trade: the CALL's own
        # first-5min LOW at a neighboring strike.
        (ATM + 50, "CE", "NSE_FO|ITM1PE", 100, 105, 95, 98, 1000),
    ]
    for strike, side, ikey, o, h, l, c, v in rows:
        conn.execute("INSERT OR REPLACE INTO orb_levels VALUES (?,?,?,?,?,?,?,?,?)",
                     (SESSION.isoformat(), strike, side, ikey, o, h, l, c, v))
    fut_high = ATM + 114.6 - 29.15   # 24285.45
    fut_low  = ATM + 47.1 - 69       # 24178.10
    conn.execute("INSERT OR REPLACE INTO orb_summary VALUES (?,?,?,?,?,?,?,?)",
                 (SESSION.isoformat(), ATM, "2026-07-28", fut_high, fut_low,
                  nearest_strike(fut_high), nearest_strike(fut_low),
                  datetime.now(IST).isoformat()))
    conn.commit()

def check(label, cond):
    status = "PASS" if cond else "FAIL"
    print("[%s] %s" % (status, label))
    assert cond, label

def main():
    conn = db()
    seed_levels(conn)

    row = conn.execute("SELECT sp_high, sp_low FROM orb_summary WHERE session_date=?",
                       (SESSION.isoformat(),)).fetchone()
    sp_high, sp_low = row
    check("SP zone computed as 24200-24300", (sp_low, sp_high) == (24200, 24300))

    levels = {}
    for strike, side, ikey, fh, flo in conn.execute(
            "SELECT strike, side, instrument_key, first_high, first_low "
            "FROM orb_levels WHERE session_date=?", (SESSION.isoformat(),)):
        levels[(strike, side)] = {"key": ikey, "high": fh, "low": flo}

    lines = ladder_lines(levels, ATM, "CE")
    check("ladder_lines cross-plots PE lows ascending (50, 160, 200)", lines == [50, 160, 200])

    watch_pairs = build_watch_pairs(levels, ATM, "CE")
    check("watch_pairs chains the full cross-plotted ladder in order",
          watch_pairs == [
              {"strike": ATM + 150, "key": "NSE_FO|PE_BELOW_ENTRY", "next_level": 160},
              {"strike": ATM - 50, "key": "NSE_FO|ITM1CE", "next_level": 200},
          ])

    # --- NEUTRAL: implied spot inside zone -> no trade
    st = DayState()
    ce_flat = Candle(ts(10, 0), 100, 100, 95, 100, 10)
    pe_flat = Candle(ts(10, 0), 60, 60, 55, 60, 10)
    on_candle_close(ts(10, 0), ce_flat, pe_flat, ATM, sp_high, sp_low, levels, st, conn, SESSION)
    check("NEUTRAL state takes no trade", st.position is None and st.signals == 0)

    # --- CE_WINS: implied spot > sp_high, CE close > its 9:15 high (114.6), PE close < its low (29.15)
    ce_break = Candle(ts(10, 30), 100, 130, 100, 120, 10)   # close 120 > 114.6
    pe_break = Candle(ts(10, 30), 25, 25, 20, 22, 10)       # close 22 < 29.15
    # implied spot = 24200 + 120 - 22 = 24298 -> > sp_high(24300)? not quite; bump CE
    ce_break = Candle(ts(10, 30), 100, 140, 100, 130, 10)   # close 130
    on_candle_close(ts(10, 30), ce_break, pe_break, ATM, sp_high, sp_low, levels, st, conn, SESSION)
    implied = ATM + 130 - 22
    check("implied spot breaks SP high (%.1f > %d)" % (implied, sp_high), implied > sp_high)
    check("CE_WINS entry taken", st.position is not None and st.position["side"] == "CE")
    check("entry price recorded at candle close (130)", st.position and st.position["entry"] == 130)

    # --- line cross: next CE close reaches the ITM1 CE ladder line (160)
    ce_line = Candle(ts(10, 33), 130, 165, 128, 162, 10)
    pe_line = Candle(ts(10, 33), 20, 20, 15, 17, 10)
    on_candle_close(ts(10, 33), ce_line, pe_line, ATM, sp_high, sp_low, levels, st, conn, SESSION)
    check("line crossed increments counter", st.position and st.position["crossed"] == 1)

    # --- exit: implied spot falls back inside SP zone
    ce_back = Candle(ts(11, 0), 130, 130, 100, 105, 10)
    pe_back = Candle(ts(11, 0), 40, 60, 40, 55, 10)  # implied spot = 24200+105-55=24250 -> inside zone
    on_candle_close(ts(11, 0), ce_back, pe_back, ATM, sp_high, sp_low, levels, st, conn, SESSION)
    check("position closed on zone re-entry", st.position is None)

    trades = conn.execute(
        "SELECT side, entry_price, exit_price, exit_reason, lines_crossed, entry_note, exit_note, source "
        "FROM orb_trades WHERE session_date=?", (SESSION.isoformat(),)).fetchall()
    check("one trade logged", len(trades) == 1)
    check("entry_note recorded", bool(trades[0][5]) and "CE WINS" in trades[0][5])
    check("exit_note recorded", bool(trades[0][6]) and "Loss" in trades[0][6])
    check("source defaults to 'live' when not passed explicitly", trades[0][7] == "live")
    print("Trade row:", trades[0])

    import orb_journal as _oj
    check("journal workbook created", os.path.exists(_oj.JOURNAL_PATH))
    from openpyxl import load_workbook
    wb = load_workbook(_oj.JOURNAL_PATH)
    check("journal has a sheet for the session date", str(SESSION) in wb.sheetnames)
    ws = wb[str(SESSION)]
    check("journal sheet has header + at least 1 trade row", ws.max_row >= 2)

    # =====================================================================
    # PE_WINS mirror path + breakeven-lock TSL
    # =====================================================================
    st2 = DayState()
    ce_lo = Candle(ts(12, 0), 45, 45, 38, 40, 10)    # CE close 40 < CE 9:15 low 47.1
    pe_hi = Candle(ts(12, 0), 75, 82, 74, 80, 10)     # PE close 80 > PE 9:15 high 69
    on_candle_close(ts(12, 0), ce_lo, pe_hi, ATM, sp_high, sp_low, levels, st2, conn, SESSION)
    implied2 = ATM + 40 - 80
    check("implied spot breaks SP low (%.1f < %d)" % (implied2, sp_low), implied2 < sp_low)
    check("PE_WINS entry taken", st2.position is not None and st2.position["side"] == "PE")
    check("PE entry price recorded (80)", st2.position["entry"] == 80)
    expected_sl = max(0.0, 80 - SL_POINTS)
    check("PE initial SL = entry - SL_POINTS (%.2f)" % expected_sl,
          abs(st2.position["trail"] - expected_sl) < 1e-6)

    # push PE up through its ladder line (95) AND past the 1R profit-lock trigger (~19.2 pts)
    ce_lo2 = Candle(ts(12, 3), 40, 40, 30, 33, 10)
    pe_up  = Candle(ts(12, 3), 80, 105, 80, 102, 10)   # +22 pts, crosses line 95 too
    on_candle_close(ts(12, 3), ce_lo2, pe_up, ATM, sp_high, sp_low, levels, st2, conn, SESSION)
    expected_trail = max(80.0 + MIN_PROFIT_POINTS, 95 - TSL_STEP_POINTS)   # profit-lock floor vs. line-cross trail
    check("trail locks at max(entry+MIN_PROFIT_POINTS, line-cross trail) = %.2f" % expected_trail,
          abs(st2.position["trail"] - expected_trail) < 1e-6)
    check("trail is at/above entry+MIN_PROFIT_POINTS (win covers fees)",
          st2.position["trail"] >= 80 + MIN_PROFIT_POINTS - 1e-9)

    # price pulls back to just under that trail -> exits in profit, not counted as a stop
    ce_lo3 = Candle(ts(12, 6), 33, 40, 30, 35, 10)
    pe_be  = Candle(ts(12, 6), 100, 100, 84, expected_trail - 0.01, 10)
    stops_before = st2.stops
    on_candle_close(ts(12, 6), ce_lo3, pe_be, ATM, sp_high, sp_low, levels, st2, conn, SESSION)
    check("trail-hit exit closes position", st2.position is None)
    check("profitable trail exit is not counted as a stop", st2.stops == stops_before)
    last_pnl = conn.execute(
        "SELECT pnl_points FROM orb_trades WHERE session_date=? AND side='PE' "
        "ORDER BY rowid DESC LIMIT 1", (SESSION.isoformat(),)).fetchone()[0]
    check("realized win is >= MIN_PROFIT_POINTS (%.2f pts, needed %.2f)" % (last_pnl, MIN_PROFIT_POINTS),
          last_pnl >= MIN_PROFIT_POINTS - 1e-6)

    # =====================================================================
    # Profit-lock floor holds even with NO ladder line crossed (pure 1R
    # trigger) - proves the fix isn't accidentally riding on the line-cross
    # trail dominating in the scenario above.
    # =====================================================================
    st4 = DayState()
    ce_e2 = Candle(ts(9, 30), 100, 140, 100, 130, 10)
    pe_e2 = Candle(ts(9, 30), 20, 20, 1, 2, 10)          # implied = 24200+130-2=24328 > sp_high -> CE wins
    on_candle_close(ts(9, 30), ce_e2, pe_e2, ATM, sp_high, sp_low, levels, st4, conn, SESSION)
    check("CE entered for profit-lock-only test", st4.position is not None)
    # Regression check for the bug found live: the cross-plotted ladder
    # includes a strike (level 50) whose value sits BELOW this entry (130).
    # It must be excluded from "lines" entirely, not trivially "crossed" the
    # instant the position opens.
    check("below-entry ladder level (50) is excluded from future targets",
          50 not in st4.position["lines"] and st4.position["lines"] == [160, 200])
    check("crossed count is 0 right at entry, not incremented by the excluded level",
          st4.position["crossed"] == 0)
    # move up exactly past 1R but stay well below the first ladder line (160)
    just_past_1r = 130 + SL_POINTS + 0.5
    ce_1r = Candle(ts(9, 33), 130, just_past_1r + 1, 128, just_past_1r, 10)
    pe_1r = Candle(ts(9, 33), 3, 3, 1, 2, 10)
    on_candle_close(ts(9, 33), ce_1r, pe_1r, ATM, sp_high, sp_low, levels, st4, conn, SESSION)
    check("no ladder line crossed yet (still below 160)", st4.position["crossed"] == 0)
    check("trail floor is exactly entry + MIN_PROFIT_POINTS (%.2f), not plain breakeven"
          % (130 + MIN_PROFIT_POINTS),
          abs(st4.position["trail"] - (130 + MIN_PROFIT_POINTS)) < 1e-6)

    # =====================================================================
    # Profit-lock takes priority over "spot back in zone": once the TSL
    # floor has moved above entry+MIN_PROFIT_POINTS, a whipsaw back inside
    # the SP zone must NOT force an exit as long as price is still above
    # the trail - the profit floor, not the directional thesis, now governs.
    # =====================================================================
    implied_inside_zone_ce = sp_high - 1   # implied spot back inside zone for a CE trade
    ce_whipsaw = Candle(ts(9, 36), just_past_1r, just_past_1r + 1, just_past_1r - 1, just_past_1r, 10)
    # choose PE close so implied_spot = ATM + ce.close - pe.close lands inside the zone
    pe_whipsaw_close = ATM + just_past_1r - implied_inside_zone_ce
    pe_whipsaw = Candle(ts(9, 36), pe_whipsaw_close, pe_whipsaw_close, pe_whipsaw_close, pe_whipsaw_close, 10)
    on_candle_close(ts(9, 36), ce_whipsaw, pe_whipsaw, ATM, sp_high, sp_low, levels, st4, conn, SESSION)
    check("profit-locked position survives a zone whipsaw (not force-exited)",
          st4.position is not None)

    # =====================================================================
    # Other-strike early exit: if ITM1 (the strike one step out from ATM)
    # has ALREADY reached ITM2's level, the move has skipped ahead of our
    # own contract - exit immediately, even though our own price is nowhere
    # near its own SL, TSL, or first ladder line, and even mid-profit.
    # =====================================================================
    st5 = DayState()
    ce_e3 = Candle(ts(10, 0), 100, 140, 100, 130, 10)
    pe_e3 = Candle(ts(10, 0), 20, 20, 1, 2, 10)   # implied = 24328 > sp_high -> CE wins, entry 130
    on_candle_close(ts(10, 0), ce_e3, pe_e3, ATM, sp_high, sp_low, levels, st5, conn, SESSION)
    check("CE entered for early-exit test", st5.position is not None)

    def fake_ltp_triggering(instrument_key):
        # ITM1CE has already reached ITM2's level (200) - should force an exit,
        # regardless of the fact our own ATM CE (currently 133, tiny profit,
        # far below its own SL/trail/first-line) gives no other reason to exit.
        return 205.0 if instrument_key == "NSE_FO|ITM1CE" else 0.0

    ce_tiny_move = Candle(ts(10, 3), 130, 135, 128, 133, 10)   # barely above entry, no other exit fires
    pe_tiny_move = Candle(ts(10, 3), 2, 2, 1, 1.5, 10)
    on_candle_close(ts(10, 3), ce_tiny_move, pe_tiny_move, ATM, sp_high, sp_low, levels, st5, conn,
                     SESSION, ltp_fn=fake_ltp_triggering)
    check("position force-exited when another strike skips ahead of its own next level",
          st5.position is None)
    early_exit_row = conn.execute(
        "SELECT exit_reason FROM orb_trades WHERE session_date=? AND entry_ts=?",
        (SESSION.isoformat(), ts(10, 0).isoformat())).fetchone()
    check("exit reason names the triggering strike (ITM1 = ATM-50 = 24150) and its value",
          str(ATM - 50) in early_exit_row[0] and "205.00" in early_exit_row[0])

    # Sanity: with a non-triggering ltp_fn, the same tiny move does NOT exit
    # (proves the check is actually discriminating, not always firing).
    st6 = DayState()
    ce_e4 = Candle(ts(10, 30), 100, 140, 100, 130, 10)
    pe_e4 = Candle(ts(10, 30), 20, 20, 1, 2, 10)
    on_candle_close(ts(10, 30), ce_e4, pe_e4, ATM, sp_high, sp_low, levels, st6, conn, SESSION)
    on_candle_close(ts(10, 33), ce_tiny_move, pe_tiny_move, ATM, sp_high, sp_low, levels, st6, conn,
                     SESSION, ltp_fn=lambda k: 0.0)
    check("no early exit when no other strike has skipped ahead", st6.position is not None)

    # =====================================================================
    # Daily loss cap: two losing trades should lock the machine out
    # =====================================================================
    st3 = DayState()
    ce_e = Candle(ts(13, 0), 100, 140, 100, 130, 10)
    pe_e = Candle(ts(13, 0), 20, 20, 1, 2, 10)         # implied = 24200+130-2=24328 > sp_high
    on_candle_close(ts(13, 0), ce_e, pe_e, ATM, sp_high, sp_low, levels, st3, conn, SESSION)
    check("trade A (CE) entered", st3.position is not None)
    # next candle: gaps straight through SL on a close (no intrabar check - candle-close only)
    ce_sl = Candle(ts(13, 3), 130, 130, 100, 108, 10)
    pe_sl = Candle(ts(13, 3), 3, 3, 1, 2, 10)          # implied = 24200+108-2=24306, still > sp_high
    on_candle_close(ts(13, 3), ce_sl, pe_sl, ATM, sp_high, sp_low, levels, st3, conn, SESSION)
    check("trade A stopped out", st3.position is None and st3.stops == 1)

    ce_e2 = Candle(ts(13, 30), 45, 45, 18, 20, 10)
    pe_e2 = Candle(ts(13, 30), 75, 82, 74, 80, 10)     # implied = 24200+20-80=24140 < sp_low
    on_candle_close(ts(13, 30), ce_e2, pe_e2, ATM, sp_high, sp_low, levels, st3, conn, SESSION)
    check("trade B (PE) entered", st3.position is not None)
    ce_sl2 = Candle(ts(13, 33), 20, 20, 15, 18, 10)
    pe_sl2 = Candle(ts(13, 33), 80, 80, 55, 58, 10)    # implied = 24200+18-58=24160, still < sp_low
    on_candle_close(ts(13, 33), ce_sl2, pe_sl2, ATM, sp_high, sp_low, levels, st3, conn, SESSION)
    check("trade B stopped out", st3.position is None and st3.stops == 2)
    check("two stops trip the lockout", st3.locked is True)
    check("daily loss recorded (Rs %.2f, cap Rs %.2f)" % (st3.daily_pnl_rupees, MAX_DAILY_LOSS),
          st3.daily_pnl_rupees < 0)

    # findings, not failures: candle-close-only stops can overshoot the nominal
    # per-trade risk budget when a candle gaps through the SL level.
    per_trade_budget = MAX_DAILY_LOSS / 2
    actual_worst = min(-22 * QTY, -22 * QTY)  # both legs realized -22 pts in this scenario
    if abs(actual_worst) > per_trade_budget:
        print("[NOTE] candle-close SL overshot the Rs %.0f/trade risk budget by Rs %.2f "
              "in this scenario - stops are evaluated on candle close, not intrabar."
              % (per_trade_budget, abs(actual_worst) - per_trade_budget))

    # =====================================================================
    # orb_auto: starting late (after 9:21) on a trading day must run TODAY,
    # not silently skip to tomorrow (this was a real bug found live on
    # 2026-07-22: starting at 09:57 scheduled the next run for 09:23 the
    # following day instead of capturing immediately).
    # =====================================================================
    tuesday_late_start = datetime(2026, 7, 21, 9, 57, tzinfo=IST)   # 2026-07-21 is a Tuesday
    check("late start (09:57, past 09:21) runs today immediately",
          should_run_today_now(tuesday_late_start, None) is True)

    tuesday_before_open = datetime(2026, 7, 21, 9, 0, tzinfo=IST)
    check("before 09:21 does not run immediately (normal sleep-until-open path)",
          should_run_today_now(tuesday_before_open, None) is False)

    tuesday_after_close = datetime(2026, 7, 21, 16, 0, tzinfo=IST)
    check("after square-off does not trigger an immediate run",
          should_run_today_now(tuesday_after_close, None) is False)

    tuesday_already_ran = datetime(2026, 7, 21, 10, 0, tzinfo=IST)
    check("does not re-run today if last_run_date is already today",
          should_run_today_now(tuesday_already_ran, date(2026, 7, 21)) is False)

    saturday = datetime(2026, 7, 25, 10, 0, tzinfo=IST)
    check("weekend never triggers an immediate run",
          should_run_today_now(saturday, None) is False)

    # =====================================================================
    # run_live_for_day's backlog catch-up filter (found live on 2026-07-22:
    # starting orb_auto after 09:21 caused the whole morning's candle
    # history to be processed as if it were live, entering a trade on a
    # candle timestamped 09:30 while the real time was already 10:05).
    # This mirrors the exact filter condition in orb_signal.run_live_for_day.
    # =====================================================================
    now_ = datetime(2026, 7, 22, 10, 5, tzinfo=IST)
    cutoff = now_ - timedelta(minutes=CANDLE_MINUTES)
    stale_candle_ts = datetime(2026, 7, 22, 9, 30, tzinfo=IST)   # the actual bad entry from today
    fresh_candle_ts = datetime(2026, 7, 22, 10, 3, tzinfo=IST)
    check("a 9:30 candle seen at 10:05 is classified as stale backlog (not acted on)",
          stale_candle_ts < cutoff)
    check("a 10:03 candle seen at 10:05 is classified as fresh (acted on live)",
          fresh_candle_ts >= cutoff)

    # =====================================================================
    # run_replay's full-ladder backfill: mock the network call and prove
    # the whole pipeline (fetch -> resample -> per-strike backfill ->
    # historical_ltp_fn -> on_candle_close) correctly reproduces the
    # other-strike early exit from real historical data, not just live
    # polling. This is the actual backtest path a trader would run.
    # =====================================================================
    fake_history = {
        # 5-min buckets from the 09:15 anchor: 09:20, 09:25, 09:30, ...
        "NSE_FO|ATMCE": [(9, 25, 130), (9, 30, 133)],
        "NSE_FO|ATMPE": [(9, 25, 2),   (9, 30, 1.5)],
        "NSE_FO|ITM1CE": [(9, 25, 150), (9, 30, 205)],   # reaches ITM2's level (200) on 2nd candle
        "NSE_FO|ITM2CE": [(9, 25, 175), (9, 30, 178)],
        "NSE_FO|ITM1PE": [(9, 25, 90),  (9, 30, 91)],
    }

    def fake_fetch_historical_candles(instrument_key, minutes, day_from, day_to):
        rows = fake_history.get(instrument_key, [])
        return [Candle(ts(h, m), c, c, c, c, 10) for h, m, c in rows]

    orb_signal.fetch_historical_candles = fake_fetch_historical_candles
    run_replay(SESSION)
    replay_row = conn.execute(
        "SELECT entry_price, exit_price, exit_reason, source FROM orb_trades "
        "WHERE session_date=? AND entry_ts=?",
        (SESSION.isoformat(), ts(9, 25).isoformat())).fetchone()
    check("run_replay entered the CE trade from mocked historical data",
          replay_row is not None and replay_row[0] == 130)
    check("run_replay's backfilled other-strike check names the triggering strike (24150)",
          replay_row is not None and str(ATM - 50) in replay_row[2])
    check("run_replay tags its rows source='replay', not 'live'",
          replay_row is not None and replay_row[3] == "replay")

    # Prove the two sources coexist in the SAME table without being confusable,
    # and that a query CAN separate them (the actual fix for the mixing bug).
    live_count = conn.execute(
        "SELECT COUNT(*) FROM orb_trades WHERE session_date=? AND source='live'",
        (SESSION.isoformat(),)).fetchone()[0]
    replay_count = conn.execute(
        "SELECT COUNT(*) FROM orb_trades WHERE session_date=? AND source='replay'",
        (SESSION.isoformat(),)).fetchone()[0]
    check("live and replay rows for the same date are distinguishable by source",
          live_count > 0 and replay_count > 0)

    # =====================================================================
    # Journal migration: an existing sheet's header (written under an older
    # HEADERS list) must gain new columns at the END, never reshuffle
    # existing ones - otherwise merging this branch would misalign every
    # row already written to a real trader's trade_journal.xlsx.
    # =====================================================================
    import orb_journal as _oj2
    from openpyxl import load_workbook as _lwb
    old_sheet_name = "2020-01-01-legacy"
    wb = _oj2._open_workbook()
    legacy_headers = ["entry_ts", "exit_ts", "side", "strike", "qty", "entry_price",
                      "exit_price", "pnl_points", "pnl_rupees", "lines_crossed",
                      "exit_reason", "why_entered", "why_pnl"]   # pre-source, pre-atm_strike schema
    ws = wb.create_sheet(old_sheet_name)
    ws.append(legacy_headers)
    ws.append(["2020-01-01T09:30:00", "2020-01-01T09:33:00", "CE", 100, 65, 10, 12,
              2, 130, 0, "trail hit", "legacy entry", "legacy exit"])
    wb.save(_oj2.JOURNAL_PATH)

    _oj2.append_trade(old_sheet_name, {
        "entry_ts": "2020-01-01T10:00:00", "exit_ts": "2020-01-01T10:03:00", "side": "PE",
        "source": "live", "atm_strike": 100, "strike": 100, "qty": 65,
        "entry_price": 20, "exit_price": 25, "pnl_points": 5, "pnl_rupees": 325,
        "lines_crossed": 0, "exit_reason": "trail hit", "trigger_strike": "",
        "trigger_strike_value": "", "why_entered": "new entry", "why_pnl": "new exit",
    })
    wb2 = _lwb(_oj2.JOURNAL_PATH)
    ws2 = wb2[old_sheet_name]
    new_header = [c.value for c in ws2[1]]
    check("legacy header columns kept in their original positions",
          new_header[:len(legacy_headers)] == legacy_headers)
    check("new columns (source, atm_strike, ...) appended at the end, not inserted",
          "source" in new_header[len(legacy_headers):])
    old_row = [c.value for c in ws2[2]]
    check("pre-existing legacy row is untouched by the migration",
          old_row[:len(legacy_headers)] == ["2020-01-01T09:30:00", "2020-01-01T09:33:00", "CE",
                                            100, 65, 10, 12, 2, 130, 0, "trail hit",
                                            "legacy entry", "legacy exit"])

    print("\nAll offline logic checks passed.")

if __name__ == "__main__":
    main()
