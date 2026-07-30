"""Tests for application.trading_analysis_dashboard."""

from __future__ import annotations

import uuid
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from application.replay_configuration import ReplayConfiguration
from application.replay_result import ReplayResult
from application.replay_session import ReplaySession, ReplayStatistics, ReplayStatus
from application.trading_analysis_dashboard import (
    CSV_HEADER,
    ChartPoint,
    ChartSeries,
    TradingAnalysisDashboard,
    TradingAnalysisFilter,
    TradingAnalysisRow,
    TradingAnalysisSummary,
    build_chart_series,
    build_rows,
    build_summary,
    filter_rows,
    write_csv,
    write_excel,
)
from business.business_result import BusinessResult
from business.pipeline_context import PipelineContext
from core.exceptions import ValidationError
from models.strike import StrikeSelection
from models.weekly_future import WeeklyFuture

_SESSION_ID = uuid.uuid4()
_DAY_1 = datetime(2026, 7, 21, 9, 20, 0, tzinfo=UTC)
_DAY_2 = datetime(2026, 7, 22, 9, 20, 0, tzinfo=UTC)
_DAY_3 = datetime(2026, 7, 23, 9, 20, 0, tzinfo=UTC)


def _context(
    timestamp: datetime,
    *,
    atm: Decimal | None = None,
    high: str | None = None,
    low: str | None = None,
    top: int | None = None,
    bottom: int | None = None,
) -> PipelineContext:
    context = PipelineContext(session_id=_SESSION_ID, candle_timestamp=timestamp)
    if atm is not None:
        context = context.with_reference_strike(atm)
    if high is not None and low is not None:
        context = context.with_weekly_future(
            WeeklyFuture(
                weekly_future_id=uuid.uuid4(),
                session_id=_SESSION_ID,
                high=Decimal(high),
                low=Decimal(low),
                calculated_at=timestamp,
            )
        )
    if top is not None and bottom is not None:
        context = context.with_selected_strike(
            StrikeSelection(
                session_id=_SESSION_ID,
                top_strike=Decimal(top),
                bottom_strike=Decimal(bottom),
                selected_at=timestamp,
            )
        )
    return context


def _replay_result(business_results: tuple[BusinessResult, ...]) -> ReplayResult:
    session = ReplaySession(
        session_id=_SESSION_ID,
        dataset_id="ds-1",
        started_at=_DAY_1,
        status=ReplayStatus.COMPLETED,
        ended_at=_DAY_3,
        statistics=ReplayStatistics(
            candles_processed=len(business_results),
            events_recorded=0,
            business_runs_succeeded=sum(1 for r in business_results if r.success),
            business_runs_failed=sum(1 for r in business_results if not r.success),
        ),
    )
    return ReplayResult(
        session=session,
        configuration=ReplayConfiguration(start_date=date(2026, 7, 21), end_date=date(2026, 7, 23)),
        generated_at=_DAY_3,
        business_results=business_results,
    )


def _three_day_result() -> ReplayResult:
    results = (
        BusinessResult(
            success=True,
            context=_context(
                _DAY_1, atm=Decimal(24200), high="24271.05", low="24198.45", top=24250, bottom=24200
            ),
        ),
        BusinessResult(
            success=True,
            context=_context(
                _DAY_2, atm=Decimal(24150), high="24144.45", low="24050.85", top=24150, bottom=24050
            ),
        ),
        BusinessResult(
            success=True,
            context=_context(
                _DAY_3, atm=Decimal(23900), high="23933.30", low="23864.85", top=23950, bottom=23850
            ),
        ),
    )
    return _replay_result(results)


class TestBuildRows:
    def test_one_row_per_day_with_computed_diffs(self) -> None:
        rows = build_rows(_three_day_result())

        assert len(rows) == 3
        row = rows[0]
        assert row.trading_day == date(2026, 7, 21)
        assert row.atm_strike == Decimal(24200)
        assert row.weekly_future_high == Decimal("24271.05")
        assert row.diff_atm_to_high == Decimal("24271.05") - Decimal(24200)
        assert row.diff_atm_to_low == Decimal("24198.45") - Decimal(24200)

    def test_range_property(self) -> None:
        rows = build_rows(_three_day_result())

        assert rows[0].range == Decimal("24271.05") - Decimal("24198.45")

    def test_missing_values_yield_none_diffs_and_range(self) -> None:
        result = _replay_result((BusinessResult(success=True, context=_context(_DAY_1)),))

        rows = build_rows(result)

        row = rows[0]
        assert row.atm_strike is None
        assert row.diff_atm_to_high is None
        assert row.diff_atm_to_low is None
        assert row.range is None

    def test_empty_result_yields_no_rows(self) -> None:
        rows = build_rows(_replay_result(()))

        assert rows == ()


class TestTradingAnalysisRowValidation:
    def test_rejects_none_trading_day(self) -> None:
        with pytest.raises(ValidationError, match="trading_day must not be None"):
            TradingAnalysisRow(
                trading_day=None,  # type: ignore[arg-type]
                atm_strike=None,
                weekly_future_high=None,
                weekly_future_low=None,
                top_strike=None,
                bottom_strike=None,
                diff_atm_to_high=None,
                diff_atm_to_low=None,
            )


class TestBuildSummary:
    def test_summary_over_three_days(self) -> None:
        rows = build_rows(_three_day_result())

        summary = build_summary(rows)

        assert summary.trading_days == 3
        expected_ranges = [row.range for row in rows]
        assert summary.average_range == sum(expected_ranges, Decimal(0)) / 3
        assert summary.largest_range == max(expected_ranges)
        assert summary.smallest_range == min(expected_ranges)
        assert (
            summary.average_distance_to_top
            == sum((abs(row.top_strike - row.atm_strike) for row in rows), Decimal(0)) / 3
        )
        assert (
            summary.average_distance_to_bottom
            == sum((abs(row.bottom_strike - row.atm_strike) for row in rows), Decimal(0)) / 3
        )

    def test_summary_of_empty_rows(self) -> None:
        summary = build_summary(())

        assert summary.trading_days == 0
        assert summary.average_range is None
        assert summary.largest_range is None
        assert summary.smallest_range is None
        assert summary.average_distance_to_top is None
        assert summary.average_distance_to_bottom is None

    def test_summary_skips_rows_missing_needed_values(self) -> None:
        result = _replay_result((BusinessResult(success=True, context=_context(_DAY_1)),))
        rows = build_rows(result)

        summary = build_summary(rows)

        assert summary.trading_days == 1
        assert summary.average_range is None
        assert summary.average_distance_to_top is None


class TestTradingAnalysisSummaryValidation:
    def test_rejects_negative_trading_days(self) -> None:
        with pytest.raises(ValidationError, match="must not be negative"):
            TradingAnalysisSummary(
                trading_days=-1,
                average_range=None,
                largest_range=None,
                smallest_range=None,
                average_distance_to_top=None,
                average_distance_to_bottom=None,
            )


class TestBuildChartSeries:
    def test_produces_five_named_series(self) -> None:
        rows = build_rows(_three_day_result())

        series = build_chart_series(rows)

        names = [s.name for s in series]
        assert names == [
            "Weekly Future High by Date",
            "Weekly Future Low by Date",
            "Top Strike by Date",
            "Bottom Strike by Date",
            "Weekly Future Range by Date",
        ]

    def test_each_series_has_one_point_per_row(self) -> None:
        rows = build_rows(_three_day_result())

        series = build_chart_series(rows)

        for s in series:
            assert len(s.points) == 3
            assert [p.trading_day for p in s.points] == [row.trading_day for row in rows]

    def test_high_series_values_match_rows(self) -> None:
        rows = build_rows(_three_day_result())

        series = build_chart_series(rows)
        high_series = series[0]

        assert [p.value for p in high_series.points] == [row.weekly_future_high for row in rows]

    def test_empty_rows_yield_empty_series(self) -> None:
        series = build_chart_series(())

        assert all(s.points == () for s in series)


class TestFilterRows:
    def test_filters_by_date_range(self) -> None:
        rows = build_rows(_three_day_result())

        filtered = filter_rows(
            rows,
            TradingAnalysisFilter(start_date=date(2026, 7, 22), end_date=date(2026, 7, 23)),
        )

        assert [row.trading_day for row in filtered] == [date(2026, 7, 22), date(2026, 7, 23)]

    def test_end_date_excludes_later_rows(self) -> None:
        rows = build_rows(_three_day_result())

        filtered = filter_rows(rows, TradingAnalysisFilter(end_date=date(2026, 7, 22)))

        assert [row.trading_day for row in filtered] == [date(2026, 7, 21), date(2026, 7, 22)]

    def test_max_range_excludes_larger_rows(self) -> None:
        rows = build_rows(_three_day_result())
        smallest = min(row.range for row in rows if row.range is not None)

        filtered = filter_rows(rows, TradingAnalysisFilter(max_range=smallest))

        assert len(filtered) == 1
        assert filtered[0].range == smallest

    def test_filters_by_atm_strike(self) -> None:
        rows = build_rows(_three_day_result())

        filtered = filter_rows(rows, TradingAnalysisFilter(atm_strike=Decimal(24150)))

        assert len(filtered) == 1
        assert filtered[0].trading_day == date(2026, 7, 22)

    def test_filters_by_range_bounds(self) -> None:
        rows = build_rows(_three_day_result())
        target_range = rows[1].range
        assert target_range is not None

        filtered = filter_rows(
            rows, TradingAnalysisFilter(min_range=target_range, max_range=target_range)
        )

        assert len(filtered) == 1
        assert filtered[0].trading_day == date(2026, 7, 22)

    def test_no_criteria_returns_all_rows(self) -> None:
        rows = build_rows(_three_day_result())

        filtered = filter_rows(rows, TradingAnalysisFilter())

        assert filtered == rows

    def test_range_filter_excludes_rows_with_no_range(self) -> None:
        result = _replay_result((BusinessResult(success=True, context=_context(_DAY_1)),))
        rows = build_rows(result)

        filtered = filter_rows(rows, TradingAnalysisFilter(min_range=Decimal(0)))

        assert filtered == ()


class TestWriteCsv:
    def test_writes_header_and_rows(self, tmp_path: Path) -> None:
        rows = build_rows(_three_day_result())
        csv_path = tmp_path / "analysis.csv"

        write_csv(rows, csv_path)

        lines = csv_path.read_text(encoding="utf-8").splitlines()
        assert lines[0] == ",".join(CSV_HEADER)
        assert len(lines) == 4

    def test_writes_empty_rows_as_header_only(self, tmp_path: Path) -> None:
        csv_path = tmp_path / "empty.csv"

        write_csv((), csv_path)

        lines = csv_path.read_text(encoding="utf-8").splitlines()
        assert lines == [",".join(CSV_HEADER)]


class TestWriteExcel:
    def test_writes_daily_and_summary_sheets(self, tmp_path: Path) -> None:
        rows = build_rows(_three_day_result())
        summary = build_summary(rows)
        xlsx_path = tmp_path / "analysis.xlsx"

        write_excel(rows, summary, xlsx_path)

        workbook = openpyxl.load_workbook(str(xlsx_path))
        assert workbook.sheetnames == ["Daily Analysis", "Summary"]
        daily = workbook["Daily Analysis"]
        assert daily.max_row == 4  # header + 3 rows
        summary_sheet = workbook["Summary"]
        assert summary_sheet.cell(row=2, column=1).value == "Trading Days"
        assert summary_sheet.cell(row=2, column=2).value == 3

    def test_writes_empty_rows_workbook(self, tmp_path: Path) -> None:
        summary = build_summary(())
        xlsx_path = tmp_path / "empty.xlsx"

        write_excel((), summary, xlsx_path)

        workbook = openpyxl.load_workbook(str(xlsx_path))
        daily = workbook["Daily Analysis"]
        assert daily.max_row == 1  # header only


class TestTradingAnalysisDashboard:
    def test_rows_property(self) -> None:
        dashboard = TradingAnalysisDashboard(_three_day_result())

        assert len(dashboard.rows) == 3

    def test_summary_method(self) -> None:
        dashboard = TradingAnalysisDashboard(_three_day_result())

        summary = dashboard.summary()

        assert summary.trading_days == 3

    def test_chart_series_method(self) -> None:
        dashboard = TradingAnalysisDashboard(_three_day_result())

        series = dashboard.chart_series()

        assert len(series) == 5

    def test_filtered_rows_method(self) -> None:
        dashboard = TradingAnalysisDashboard(_three_day_result())

        filtered = dashboard.filtered_rows(TradingAnalysisFilter(atm_strike=Decimal(24150)))

        assert len(filtered) == 1

    def test_export_csv_method(self, tmp_path: Path) -> None:
        dashboard = TradingAnalysisDashboard(_three_day_result())
        csv_path = tmp_path / "out.csv"

        dashboard.export_csv(csv_path)

        assert csv_path.exists()

    def test_export_excel_method(self, tmp_path: Path) -> None:
        dashboard = TradingAnalysisDashboard(_three_day_result())
        xlsx_path = tmp_path / "out.xlsx"

        dashboard.export_excel(xlsx_path)

        assert xlsx_path.exists()


class TestChartPointAndSeriesConstruction:
    def test_chart_point_holds_date_and_value(self) -> None:
        point = ChartPoint(trading_day=date(2026, 7, 21), value=Decimal(100))

        assert point.trading_day == date(2026, 7, 21)
        assert point.value == Decimal(100)

    def test_chart_series_defaults_to_empty_points(self) -> None:
        series = ChartSeries(name="Test Series")

        assert series.points == ()
