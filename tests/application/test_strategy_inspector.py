"""Tests for application.strategy_inspector."""

from __future__ import annotations

import json
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from application.replay_configuration import ReplayConfiguration
from application.replay_result import ReplayResult
from application.replay_session import ReplaySession, ReplayStatistics, ReplayStatus
from application.strategy_inspector import (
    ORBSection,
    ReferenceLevelsSection,
    StrategyInspector,
    StrategyInspectorView,
    WeeklyFutureSection,
    build_view,
    write_csv,
    write_excel,
    write_json,
)
from application.strategy_timeline import build_strategy_timeline
from business.business_result import BusinessResult
from business.pipeline_context import PipelineContext
from core.enums import OptionType, ORBStatus
from core.exceptions import ValidationError
from models.orb_result import ORBResult
from models.reference_level import ReferenceLevel
from models.strike import StrikeSelection
from models.weekly_future import WeeklyFuture

_SESSION_ID = uuid.uuid4()
_T1 = datetime(2026, 7, 29, 9, 20, 0, tzinfo=UTC)
_T2 = datetime(2026, 7, 29, 9, 25, 0, tzinfo=UTC)
_ENDED_AT = datetime(2026, 7, 29, 9, 30, 0, tzinfo=UTC)


def _reference_level() -> ReferenceLevel:
    return ReferenceLevel(
        strike=Decimal(24200),
        ce_high=Decimal("143.45"),
        ce_low=Decimal(116),
        pe_high=Decimal("165.8"),
        pe_low=Decimal(128),
    )


def _ladder() -> tuple[ReferenceLevel, ...]:
    others = tuple(
        ReferenceLevel(
            strike=Decimal(24000 + i * 50),
            ce_high=Decimal(150),
            ce_low=Decimal(100),
            pe_high=Decimal(150),
            pe_low=Decimal(100),
        )
        for i in range(13)
        if 24000 + i * 50 != 24200
    )
    return (_reference_level(), *others[:12])


def _full_context(timestamp: datetime) -> PipelineContext:
    weekly_future = WeeklyFuture(
        weekly_future_id=uuid.uuid4(),
        session_id=_SESSION_ID,
        high=Decimal("24215.45"),
        low=Decimal("24150.2"),
        calculated_at=timestamp,
    )
    selected_strike = StrikeSelection(
        session_id=_SESSION_ID,
        top_strike=Decimal(24200),
        bottom_strike=Decimal(24150),
        selected_at=timestamp,
    )
    orb_result = ORBResult(
        orb_result_id=uuid.uuid4(),
        session_id=_SESSION_ID,
        strike=Decimal(24200),
        side=OptionType.CALL,
        opening_high=Decimal("143.45"),
        opening_low=Decimal(116),
        range=Decimal("27.45"),
        status=ORBStatus.BREAKOUT,
        calculated_at=timestamp,
    )
    return (
        PipelineContext(
            session_id=_SESSION_ID, candle_timestamp=timestamp, reference_data=_ladder()
        )
        .with_reference_strike(Decimal(24200))
        .with_weekly_future(weekly_future)
        .with_selected_strike(selected_strike)
        .with_orb_result(orb_result)
    )


def _replay_result(
    business_results: tuple[BusinessResult, ...], *, status: ReplayStatus = ReplayStatus.COMPLETED
) -> ReplayResult:
    session = ReplaySession(
        session_id=_SESSION_ID,
        dataset_id="ds-1",
        started_at=_T1,
        status=status,
        ended_at=_ENDED_AT if status != ReplayStatus.RUNNING else None,
        statistics=(
            ReplayStatistics(
                candles_processed=len(business_results),
                events_recorded=0,
                business_runs_succeeded=sum(1 for r in business_results if r.success),
                business_runs_failed=sum(1 for r in business_results if not r.success),
            )
            if status != ReplayStatus.RUNNING
            else None
        ),
    )
    timeline = build_strategy_timeline(business_results, _ENDED_AT)
    return ReplayResult(
        session=session,
        configuration=ReplayConfiguration(start_date=date(2026, 7, 29), end_date=date(2026, 7, 29)),
        generated_at=_ENDED_AT,
        business_results=business_results,
        strategy_timeline=timeline,
    )


class TestBuildViewRendering:
    def test_full_result_renders_every_section(self) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)

        view = build_view(result)

        assert view.summary.trading_date == date(2026, 7, 29)
        assert view.summary.replay_status == ReplayStatus.COMPLETED
        assert view.summary.pipeline_status == "PASSED"
        assert view.reference_levels.atm == Decimal(24200)
        assert len(view.reference_levels.ladder) == 13
        assert view.weekly_future.high == Decimal("24215.45")
        assert view.weekly_future.top_strike == Decimal(24200)
        assert view.orb.opening_high == Decimal("143.45")
        assert view.orb.status == ORBStatus.BREAKOUT
        assert len(view.timeline.events) > 0

    def test_uses_last_candle_for_final_state(self) -> None:
        early = BusinessResult(success=True, context=_full_context(_T1))
        # A later, different ORB status than the first candle's.
        later_context = _full_context(_T2).with_orb_result(
            ORBResult(
                orb_result_id=uuid.uuid4(),
                session_id=_SESSION_ID,
                strike=Decimal(24200),
                side=OptionType.CALL,
                opening_high=Decimal("143.45"),
                opening_low=Decimal(116),
                range=Decimal("27.45"),
                status=ORBStatus.BREAKDOWN,
                calculated_at=_T2,
            )
        )
        later = BusinessResult(success=True, context=later_context)
        result = _replay_result((early, later))

        view = build_view(result)

        assert view.orb.status == ORBStatus.BREAKDOWN
        assert view.summary.trading_date == date(2026, 7, 29)

    def test_pipeline_status_failed_when_any_candle_failed(self) -> None:
        results = (
            BusinessResult(success=True, context=_full_context(_T1)),
            BusinessResult(success=False, context=_full_context(_T2)),
        )
        result = _replay_result(results)

        view = build_view(result)

        assert view.summary.pipeline_status == "FAILED"


class TestNullHandling:
    def test_empty_business_results_yields_all_none_sections(self) -> None:
        result = _replay_result((), status=ReplayStatus.RUNNING)

        view = build_view(result)

        assert view.summary.trading_date is None
        assert view.reference_levels.atm is None
        assert view.reference_levels.ladder == ()
        assert view.weekly_future.high is None
        assert view.weekly_future.top_strike is None
        assert view.orb.opening_high is None
        assert view.orb.status is None
        assert view.summary.pipeline_status == "FAILED"

    def test_stage_never_ran_yields_none_for_that_section_only(self) -> None:
        bare_context = PipelineContext(session_id=_SESSION_ID, candle_timestamp=_T1)
        results = (BusinessResult(success=True, context=bare_context),)
        result = _replay_result(results)

        view = build_view(result)

        assert view.summary.trading_date == date(2026, 7, 29)
        assert view.weekly_future.high is None
        assert view.orb.status is None


class TestLargeReplaySession:
    def test_handles_many_candles_correctly(self) -> None:
        results = tuple(
            BusinessResult(success=True, context=_full_context(_T1 + timedelta(minutes=i)))
            for i in range(500)
        )
        result = _replay_result(results)

        view = build_view(result)

        assert view.summary.pipeline_status == "PASSED"
        assert view.orb.status == ORBStatus.BREAKOUT
        # ORB_CALCULATED emitted once per candle - 500 candles worth.
        orb_events = [e for e in view.timeline.events if e.event_type.value == "ORB_CALCULATED"]
        assert len(orb_events) == 500


class TestTimelineOrdering:
    def test_timeline_events_are_chronological(self) -> None:
        results = (
            BusinessResult(success=True, context=_full_context(_T1)),
            BusinessResult(success=True, context=_full_context(_T2)),
        )
        result = _replay_result(results)

        view = build_view(result)

        timestamps = [event.timestamp for event in view.timeline.events]
        assert timestamps == sorted(timestamps)

    def test_timeline_is_unchanged_from_replay_result(self) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)

        view = build_view(result)

        assert view.timeline is result.strategy_timeline


class TestStrategyInspectorViewValidation:
    def test_rejects_none_summary(self) -> None:
        with pytest.raises(ValidationError, match="summary must not be None"):
            StrategyInspectorView(
                summary=None,  # type: ignore[arg-type]
                reference_levels=ReferenceLevelsSection(atm=None),
                weekly_future=WeeklyFutureSection(None, None, None, None),
                orb=ORBSection(None, None, None, None),
                timeline=build_strategy_timeline((), _ENDED_AT),
            )


class TestExportJson:
    def test_writes_all_sections(self, tmp_path: Path) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)
        view = build_view(result)
        json_path = tmp_path / "inspector.json"

        write_json(view, json_path)

        data = json.loads(json_path.read_text(encoding="utf-8"))
        assert data["replay_summary"]["pipeline_status"] == "PASSED"
        assert data["reference_levels"]["atm"] == "24200"
        assert len(data["reference_levels"]["ladder"]) == 13
        assert data["weekly_future"]["high"] == "24215.45"
        assert data["orb"]["status"] == "BREAKOUT"
        assert len(data["timeline"]) > 0

    def test_writes_none_values_for_empty_result(self, tmp_path: Path) -> None:
        result = _replay_result((), status=ReplayStatus.RUNNING)
        view = build_view(result)
        json_path = tmp_path / "empty.json"

        write_json(view, json_path)

        data = json.loads(json_path.read_text(encoding="utf-8"))
        assert data["replay_summary"]["trading_date"] is None
        assert data["weekly_future"]["high"] is None
        assert data["orb"]["status"] is None


class TestExportCsv:
    def test_writes_summary_ladder_and_timeline_rows(self, tmp_path: Path) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)
        view = build_view(result)
        csv_path = tmp_path / "inspector.csv"

        write_csv(view, csv_path)

        lines = csv_path.read_text(encoding="utf-8").splitlines()
        assert lines[0] == "section,field,value"
        assert any("reference_ladder" in line for line in lines)
        assert any("timeline" in line for line in lines)


class TestExportExcel:
    def test_writes_three_sheets(self, tmp_path: Path) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)
        view = build_view(result)
        xlsx_path = tmp_path / "inspector.xlsx"

        write_excel(view, xlsx_path)

        workbook = openpyxl.load_workbook(str(xlsx_path))
        assert workbook.sheetnames == ["Summary", "Reference Ladder", "Timeline"]
        assert workbook["Reference Ladder"].max_row == 14  # header + 13 levels

    def test_writes_empty_result_without_crashing(self, tmp_path: Path) -> None:
        result = _replay_result((), status=ReplayStatus.RUNNING)
        view = build_view(result)
        xlsx_path = tmp_path / "empty.xlsx"

        write_excel(view, xlsx_path)

        workbook = openpyxl.load_workbook(str(xlsx_path))
        assert workbook["Reference Ladder"].max_row == 1  # header only


class TestStrategyInspectorFacade:
    def test_view_property(self) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)

        inspector = StrategyInspector(result)

        assert inspector.view.summary.pipeline_status == "PASSED"

    def test_export_methods(self, tmp_path: Path) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)
        inspector = StrategyInspector(result)

        inspector.export_json(tmp_path / "out.json")
        inspector.export_csv(tmp_path / "out.csv")
        inspector.export_excel(tmp_path / "out.xlsx")

        assert (tmp_path / "out.json").exists()
        assert (tmp_path / "out.csv").exists()
        assert (tmp_path / "out.xlsx").exists()

    def test_never_mutates_the_replay_result(self, tmp_path: Path) -> None:
        results = (BusinessResult(success=True, context=_full_context(_T1)),)
        result = _replay_result(results)
        original_business_results = result.business_results

        StrategyInspector(result)

        assert result.business_results is original_business_results
