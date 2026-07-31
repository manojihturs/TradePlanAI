"""Tests for application.replay_comparison."""

from __future__ import annotations

import json
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

from application.replay_comparison import (
    FieldDifference,
    ReplayComparison,
    build_comparison,
    write_csv,
    write_excel,
    write_json,
)
from application.replay_configuration import ReplayConfiguration
from application.replay_result import ReplayResult
from application.replay_session import ReplaySession, ReplayStatistics, ReplayStatus
from application.strategy_timeline import build_strategy_timeline
from business.business_result import BusinessResult
from business.pipeline_context import PipelineContext
from core.enums import OptionType, ORBStatus
from models.orb_result import ORBResult
from models.reference_level import ReferenceLevel
from models.strike import StrikeSelection
from models.weekly_future import WeeklyFuture

_SESSION_A = uuid.uuid4()
_SESSION_B = uuid.uuid4()
_T0 = datetime(2026, 7, 29, 9, 20, 0, tzinfo=UTC)


def _reference_level(strike: int = 24200) -> ReferenceLevel:
    return ReferenceLevel(
        strike=Decimal(strike),
        ce_high=Decimal("143.45"),
        ce_low=Decimal(116),
        pe_high=Decimal("165.8"),
        pe_low=Decimal(128),
    )


def _context(
    session_id: uuid.UUID,
    timestamp: datetime,
    *,
    wf_high: str = "24215.45",
    wf_low: str = "24150.2",
    top_strike: int = 24200,
    bottom_strike: int = 24150,
    orb_status: ORBStatus = ORBStatus.NONE,
    strike: int = 24200,
) -> PipelineContext:
    weekly_future = WeeklyFuture(
        weekly_future_id=uuid.uuid4(),
        session_id=session_id,
        high=Decimal(wf_high),
        low=Decimal(wf_low),
        calculated_at=timestamp,
    )
    selected_strike = StrikeSelection(
        session_id=session_id,
        top_strike=Decimal(top_strike),
        bottom_strike=Decimal(bottom_strike),
        selected_at=timestamp,
    )
    orb_result = ORBResult(
        orb_result_id=uuid.uuid4(),
        session_id=session_id,
        strike=Decimal(strike),
        side=OptionType.CALL,
        opening_high=Decimal("143.45"),
        opening_low=Decimal(116),
        range=Decimal("27.45"),
        status=orb_status,
        calculated_at=timestamp,
    )
    return (
        PipelineContext(
            session_id=session_id,
            candle_timestamp=timestamp,
            reference_data=(_reference_level(strike),),
        )
        .with_reference_strike(Decimal(strike))
        .with_weekly_future(weekly_future)
        .with_selected_strike(selected_strike)
        .with_orb_result(orb_result)
    )


def _replay_result(
    session_id: uuid.UUID, business_results: tuple[BusinessResult, ...], dataset_id: str = "ds"
) -> ReplayResult:
    ended_at = business_results[-1].context.candle_timestamp if business_results else _T0
    session = ReplaySession(
        session_id=session_id,
        dataset_id=dataset_id,
        started_at=_T0,
        status=ReplayStatus.COMPLETED,
        ended_at=ended_at,
        statistics=ReplayStatistics(
            candles_processed=len(business_results),
            events_recorded=0,
            business_runs_succeeded=sum(1 for r in business_results if r.success),
            business_runs_failed=sum(1 for r in business_results if not r.success),
        ),
    )
    timeline = build_strategy_timeline(business_results, ended_at)
    return ReplayResult(
        session=session,
        configuration=ReplayConfiguration(start_date=date(2026, 7, 29), end_date=date(2026, 7, 29)),
        generated_at=ended_at,
        business_results=business_results,
        strategy_timeline=timeline,
    )


def _identical_pair() -> tuple[ReplayResult, ReplayResult]:
    result_a = _replay_result(
        _SESSION_A,
        (BusinessResult(success=True, context=_context(_SESSION_A, _T0)),),
        dataset_id="NIFTY-A",
    )
    result_b = _replay_result(
        _SESSION_B,
        (BusinessResult(success=True, context=_context(_SESSION_B, _T0)),),
        dataset_id="NIFTY-B",
    )
    return result_a, result_b


def _different_pair() -> tuple[ReplayResult, ReplayResult]:
    result_a = _replay_result(
        _SESSION_A,
        (
            BusinessResult(
                success=True,
                context=_context(_SESSION_A, _T0, orb_status=ORBStatus.BREAKOUT),
            ),
        ),
        dataset_id="NIFTY-A",
    )
    result_b = _replay_result(
        _SESSION_B,
        (
            BusinessResult(
                success=True,
                context=_context(
                    _SESSION_B,
                    _T0,
                    wf_high="24300",
                    wf_low="24100",
                    top_strike=24300,
                    bottom_strike=24100,
                    orb_status=ORBStatus.BREAKDOWN,
                ),
            ),
        ),
        dataset_id="NIFTY-B",
    )
    return result_a, result_b


class TestBuildComparisonIdentical:
    def test_no_differences_when_both_sessions_match(self) -> None:
        result_a, result_b = _identical_pair()

        report = build_comparison(result_a, result_b)

        assert report.has_differences is False
        assert all(not d.differs for d in report.weekly_future_differences)
        assert all(not d.differs for d in report.strike_differences)

    def test_dataset_ids_are_recorded(self) -> None:
        result_a, result_b = _identical_pair()

        report = build_comparison(result_a, result_b)

        assert report.dataset_id_a == "NIFTY-A"
        assert report.dataset_id_b == "NIFTY-B"


class TestBuildComparisonDifferences:
    def test_weekly_future_differences_detected(self) -> None:
        result_a, result_b = _different_pair()

        report = build_comparison(result_a, result_b)

        high_diff = next(d for d in report.weekly_future_differences if d.field == "high")
        assert high_diff.differs is True
        assert high_diff.value_a == "24215.45"
        assert high_diff.value_b == "24300"

    def test_strike_differences_detected(self) -> None:
        result_a, result_b = _different_pair()

        report = build_comparison(result_a, result_b)

        top_diff = next(d for d in report.strike_differences if d.field == "top_strike")
        assert top_diff.differs is True
        assert top_diff.value_a == "24200"
        assert top_diff.value_b == "24300"

    def test_orb_status_difference_detected(self) -> None:
        result_a, result_b = _different_pair()

        report = build_comparison(result_a, result_b)

        status_diff = next(d for d in report.orb_differences if d.field == "status")
        assert status_diff.differs is True
        assert status_diff.value_a == "BREAKOUT"
        assert status_diff.value_b == "BREAKDOWN"

    def test_has_differences_is_true(self) -> None:
        result_a, result_b = _different_pair()

        report = build_comparison(result_a, result_b)

        assert report.has_differences is True

    def test_never_mutates_either_replay_result(self) -> None:
        result_a, result_b = _different_pair()
        events_a = result_a.strategy_timeline.events
        events_b = result_b.strategy_timeline.events

        build_comparison(result_a, result_b)

        assert result_a.strategy_timeline.events is events_a
        assert result_b.strategy_timeline.events is events_b


class TestTimelineComparison:
    def test_total_events_and_breakout_breakdown_counts(self) -> None:
        result_a, result_b = _different_pair()

        report = build_comparison(result_a, result_b)

        assert report.timeline.total_events_a == report.timeline.total_events_b
        assert report.timeline.breakout_count_a == 1
        assert report.timeline.breakout_count_b == 0
        assert report.timeline.breakdown_count_a == 0
        assert report.timeline.breakdown_count_b == 1

    def test_stages_and_event_types_symmetric_when_identical(self) -> None:
        result_a, result_b = _identical_pair()

        report = build_comparison(result_a, result_b)

        assert report.timeline.stages_only_in_a == ()
        assert report.timeline.stages_only_in_b == ()
        assert report.timeline.event_types_only_in_a == ()
        assert report.timeline.event_types_only_in_b == ()

    def test_stages_only_in_one_side(self) -> None:
        result_a = _replay_result(
            _SESSION_A,
            (BusinessResult(success=True, context=_context(_SESSION_A, _T0)),),
        )
        result_b = _replay_result(_SESSION_B, ())  # only REPLAY_FINISHED

        report = build_comparison(result_a, result_b)

        assert "orb" in report.timeline.stages_only_in_a
        assert report.timeline.stages_only_in_b == ()


class TestDurationAndPipelineDiagnostics:
    def test_duration_and_status_compared(self) -> None:
        result_a, result_b = _identical_pair()

        report = build_comparison(result_a, result_b)

        assert len(report.duration_differences) == 2

    def test_pipeline_diagnostics_compared(self) -> None:
        result_a = _replay_result(
            _SESSION_A,
            (
                BusinessResult(success=True, context=_context(_SESSION_A, _T0)),
                BusinessResult(success=False, context=_context(_SESSION_A, _T0)),
            ),
        )
        result_b = _replay_result(
            _SESSION_B,
            (BusinessResult(success=True, context=_context(_SESSION_B, _T0)),),
        )

        report = build_comparison(result_a, result_b)

        succeeded_diff = next(
            d for d in report.pipeline_diagnostics_differences if d.field == "succeeded_count"
        )
        failed_diff = next(
            d for d in report.pipeline_diagnostics_differences if d.field == "failed_count"
        )
        assert succeeded_diff.value_a == "1"
        assert succeeded_diff.value_b == "1"
        assert failed_diff.value_a == "1"
        assert failed_diff.value_b == "0"
        assert failed_diff.differs is True


class TestFieldDifferenceHandlesNone:
    def test_both_none_does_not_differ(self) -> None:
        result_a = _replay_result(_SESSION_A, ())
        result_b = _replay_result(_SESSION_B, ())

        report = build_comparison(result_a, result_b)

        for diff in report.weekly_future_differences:
            assert diff.value_a is None
            assert diff.value_b is None
            assert diff.differs is False


class TestExportJson:
    def test_writes_all_sections(self, tmp_path: Path) -> None:
        result_a, result_b = _different_pair()
        report = build_comparison(result_a, result_b)
        json_path = tmp_path / "comparison.json"

        write_json(report, json_path)

        data = json.loads(json_path.read_text(encoding="utf-8"))
        assert data["has_differences"] is True
        assert data["dataset_id_a"] == "NIFTY-A"
        assert len(data["weekly_future"]) == 2
        assert data["timeline"]["breakout_count_a"] == 1


class TestExportCsv:
    def test_writes_section_rows_and_timeline_rows(self, tmp_path: Path) -> None:
        result_a, result_b = _different_pair()
        report = build_comparison(result_a, result_b)
        csv_path = tmp_path / "comparison.csv"

        write_csv(report, csv_path)

        lines = csv_path.read_text(encoding="utf-8").splitlines()
        assert lines[0] == "section,field,value_a,value_b,differs"
        assert any("weekly_future" in line for line in lines)
        assert any("timeline" in line for line in lines)


class TestExportExcel:
    def test_writes_two_sheets(self, tmp_path: Path) -> None:
        result_a, result_b = _different_pair()
        report = build_comparison(result_a, result_b)
        xlsx_path = tmp_path / "comparison.xlsx"

        write_excel(report, xlsx_path)

        workbook = openpyxl.load_workbook(str(xlsx_path))
        assert workbook.sheetnames == ["Differences", "Timeline"]
        assert workbook["Differences"].max_row > 1
        assert workbook["Timeline"].max_row == 8


class TestReplayComparisonFacade:
    def test_report_property(self) -> None:
        result_a, result_b = _different_pair()

        comparison = ReplayComparison(result_a, result_b)

        assert comparison.report.has_differences is True

    def test_export_methods(self, tmp_path: Path) -> None:
        result_a, result_b = _different_pair()
        comparison = ReplayComparison(result_a, result_b)

        comparison.export_json(tmp_path / "out.json")
        comparison.export_csv(tmp_path / "out.csv")
        comparison.export_excel(tmp_path / "out.xlsx")

        assert (tmp_path / "out.json").exists()
        assert (tmp_path / "out.csv").exists()
        assert (tmp_path / "out.xlsx").exists()

    def test_never_mutates_either_result(self) -> None:
        result_a, result_b = _different_pair()
        events_a = result_a.strategy_timeline.events
        events_b = result_b.strategy_timeline.events

        ReplayComparison(result_a, result_b)

        assert result_a.strategy_timeline.events is events_a
        assert result_b.strategy_timeline.events is events_b


class TestFieldDifferenceDataclass:
    def test_construction(self) -> None:
        diff = FieldDifference(field="high", value_a="1", value_b="2", differs=True)

        assert diff.field == "high"
        assert diff.differs is True
