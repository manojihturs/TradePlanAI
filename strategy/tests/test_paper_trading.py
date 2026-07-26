"""Unit tests for strategy.paper_trading (Module 6).

No network access and no broker calls anywhere in this module or
these tests - Excel export is verified by reading the file back with
openpyxl.
"""

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from strategy.entry_signal import EntrySignal, MappingAnchor, TradeSide
from strategy.exit_signal import ExitLevels, ExitReason, ExitSignal
from strategy.position_manager import ClosedTrade
from strategy.paper_trading import PaperTradingEngine, Trade, build_trade


def _row1_closed_trade_and_levels():
    """Real 2026-07-22 row 1: PE @ strike 24050, entry 206.35,
    target 238.75, stop_loss 176.85, exit 09:45 @ 238.75 (target hit)."""
    entry = EntrySignal(
        timestamp=datetime(2026, 7, 22, 9, 25), strike=24050,
        side=TradeSide.PE, anchor=MappingAnchor.TOP,
        ce_level=114.00, pe_level=206.35,
    )
    exit_signal = ExitSignal(
        timestamp=datetime(2026, 7, 22, 9, 45),
        reason=ExitReason.TARGET, exit_price=238.75,
    )
    exit_levels = ExitLevels(target=238.75, stop_loss=176.85)
    closed = ClosedTrade(entry=entry, exit=exit_signal)
    return closed, exit_levels


class BuildTradeTests(unittest.TestCase):
    def test_matches_real_row1_exactly(self) -> None:
        closed, exit_levels = _row1_closed_trade_and_levels()
        trade = build_trade(closed, exit_levels)

        self.assertEqual(trade.trade_date, datetime(2026, 7, 22).date())
        self.assertEqual(trade.trade_time, datetime(2026, 7, 22, 9, 25).time())
        self.assertEqual(trade.side, TradeSide.PE)
        self.assertEqual(trade.entry, 206.35)
        self.assertEqual(trade.exit, 238.75)
        self.assertEqual(trade.reason, "TARGET")
        self.assertEqual(trade.target, 238.75)
        self.assertEqual(trade.stop_loss, 176.85)
        self.assertAlmostEqual(trade.pnl, 32.40, places=2)

    def test_ce_side_uses_ce_level_as_entry(self) -> None:
        entry = EntrySignal(
            timestamp=datetime(2026, 7, 22, 10, 50), strike=24000,
            side=TradeSide.CE, anchor=MappingAnchor.TOP,
            ce_level=94.40, pe_level=238.75,
        )
        exit_signal = ExitSignal(
            timestamp=datetime(2026, 7, 22, 11, 5),
            reason=ExitReason.STOP_LOSS, exit_price=78.20,
        )
        exit_levels = ExitLevels(target=114.00, stop_loss=78.20)
        closed = ClosedTrade(entry=entry, exit=exit_signal)

        trade = build_trade(closed, exit_levels)
        self.assertEqual(trade.entry, 94.40)
        self.assertEqual(trade.pnl, 78.20 - 94.40)


class PaperTradingEngineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = PaperTradingEngine()

    def test_starts_with_no_trades(self) -> None:
        self.assertEqual(self.engine.trades, [])

    def test_record_trade_stores_and_returns_trade(self) -> None:
        closed, exit_levels = _row1_closed_trade_and_levels()
        trade = self.engine.record_trade(closed, exit_levels)
        self.assertIsInstance(trade, Trade)
        self.assertEqual(self.engine.trades, [trade])

    def test_trades_property_returns_a_copy(self) -> None:
        closed, exit_levels = _row1_closed_trade_and_levels()
        self.engine.record_trade(closed, exit_levels)
        snapshot = self.engine.trades
        snapshot.clear()
        self.assertEqual(len(self.engine.trades), 1)   # unaffected by mutating the copy

    def test_export_to_excel_writes_correct_rows(self) -> None:
        closed, exit_levels = _row1_closed_trade_and_levels()
        self.engine.record_trade(closed, exit_levels)

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "trades.xlsx"
            result_path = self.engine.export_to_excel(path)
            self.assertEqual(result_path, path)
            self.assertTrue(path.exists())

            wb = load_workbook(path)
            sheet = wb["Paper Trades"]
            header = [c.value for c in sheet[1]]
            self.assertEqual(header, ["Date", "Time", "Side", "Entry", "Exit",
                                       "Reason", "Target", "Stop Loss", "PnL"])
            row = [c.value for c in sheet[2]]
            self.assertEqual(row[0], "2026-07-22")
            self.assertEqual(row[2], "PE")
            self.assertEqual(row[3], 206.35)
            self.assertEqual(row[4], 238.75)
            self.assertEqual(row[5], "TARGET")
            self.assertAlmostEqual(row[8], 32.40, places=2)

    def test_export_with_zero_trades_writes_header_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.xlsx"
            self.engine.export_to_excel(path)
            wb = load_workbook(path)
            sheet = wb["Paper Trades"]
            self.assertEqual(sheet.max_row, 1)   # header only


if __name__ == "__main__":
    unittest.main()
