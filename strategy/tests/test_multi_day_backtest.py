"""Unit tests for strategy.multi_day_backtest (Module 10).

No network access required. Modules 1-9 are used unmodified.
"""

import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from strategy.entry_signal import Candle
from strategy.multi_day_backtest import (
    DayInput,
    MultiDayBacktestError,
    export_to_excel,
    run_multi_day_backtest,
)

STRIKE_GAP = 50
NUM_STRIKES = 6

# Real 2026-07-22 first-5min levels, ATM 24150.
_RAW_LEVELS = {
    23850: (342.80, 288.00, 86.50, 52.65),
    23900: (310.00, 253.80, 102.65, 63.65),
    23950: (269.95, 221.85, 120.85, 78.20),
    24000: (238.75, 192.00, 141.65, 94.40),
    24050: (206.35, 165.50, 164.65, 114.00),
    24100: (176.85, 140.75, 190.00, 135.05),
    24150: (154.00, 118.85, 218.00, 159.55),
    24200: (133.00, 99.10, 248.40, 186.05),
    24250: (106.45, 81.70, 280.65, 215.05),
    24300: (92.50, 66.80, 315.40, 245.05),
    24350: (76.00, 53.65, 352.35, 280.00),
    24400: (59.75, 43.05, 391.65, 320.35),
    24450: (49.60, 34.20, 433.20, 357.55),
}

_CE_SAFE = Candle(90.0, 95.0, 85.0, 90.0)
_PE_SAFE = Candle(150.0, 155.0, 145.0, 150.0)


def _fetcher(strike, side):
    ce_h, ce_l, pe_h, pe_l = _RAW_LEVELS[strike]
    return (0.0, ce_h if side == "CE" else pe_h, ce_l if side == "CE" else pe_l, 0.0)


def _failing_fetcher(strike, side):
    raise RuntimeError("simulated broker outage")


def _one_trade_day(session_date: date) -> DayInput:
    """A day with exactly one clean trade at strike 24050 (real row 1):
    PE entry 09:25 @ 206.35, target hit 09:45 @ 238.75."""
    strike = 24050
    ce_series = {strike: {
        datetime(session_date.year, session_date.month, session_date.day, 9, 20): _CE_SAFE,
        datetime(session_date.year, session_date.month, session_date.day, 9, 25): _CE_SAFE,
        datetime(session_date.year, session_date.month, session_date.day, 9, 45): _CE_SAFE,
    }}
    pe_series = {strike: {
        datetime(session_date.year, session_date.month, session_date.day, 9, 20): _PE_SAFE,
        datetime(session_date.year, session_date.month, session_date.day, 9, 25):
            Candle(200.0, 220.0, 200.0, 217.2),
        datetime(session_date.year, session_date.month, session_date.day, 9, 45):
            Candle(233.0, 240.0, 230.0, 236.0),
    }}
    return DayInput(session_date=session_date, spot_open=24150.0,
                     fetch_first_candle=_fetcher, ce_series=ce_series, pe_series=pe_series)


def _empty_day(session_date: date) -> DayInput:
    """A day with no crossings at all - zero trades expected."""
    strike = 24050
    ts = datetime(session_date.year, session_date.month, session_date.day, 9, 20)
    return DayInput(session_date=session_date, spot_open=24150.0,
                     fetch_first_candle=_fetcher,
                     ce_series={strike: {ts: _CE_SAFE}}, pe_series={strike: {ts: _PE_SAFE}})


class RunMultiDayBacktestTests(unittest.TestCase):
    def test_two_clean_days_produce_two_reports(self) -> None:
        days = [date(2026, 7, 22), date(2026, 7, 23)]

        def fetch_day(d):
            return _one_trade_day(d)

        result = run_multi_day_backtest(days, fetch_day, STRIKE_GAP, NUM_STRIKES)

        self.assertEqual(len(result.daily_reports), 2)
        self.assertEqual(len(result.failures), 0)
        for report in result.daily_reports:
            self.assertEqual(report.trades, 1)
            self.assertEqual(report.wins, 1)
            self.assertAlmostEqual(report.net_pnl, 32.40, places=2)
        self.assertEqual(result.overall.trading_days, 2)
        self.assertEqual(result.overall.total_trades, 2)
        self.assertAlmostEqual(result.overall.net_pnl, 64.80, places=2)

    def test_state_resets_between_days_same_level_can_trade_again(self) -> None:
        # If state leaked across days, day 2's identical trade at the SAME
        # level would be rejected as "already completed" from day 1.
        days = [date(2026, 7, 22), date(2026, 7, 23)]
        result = run_multi_day_backtest(days, lambda d: _one_trade_day(d), STRIKE_GAP, NUM_STRIKES)
        self.assertEqual(result.daily_reports[0].trades, 1)
        self.assertEqual(result.daily_reports[1].trades, 1)   # not blocked by day 1's USED level

    def test_failing_day_is_recorded_not_raised_and_others_continue(self) -> None:
        days = [date(2026, 7, 22), date(2026, 7, 23), date(2026, 7, 24)]

        def fetch_day(d):
            if d == date(2026, 7, 23):
                raise RuntimeError("simulated fetch failure for this day")
            return _one_trade_day(d)

        result = run_multi_day_backtest(days, fetch_day, STRIKE_GAP, NUM_STRIKES)

        self.assertEqual(len(result.daily_reports), 2)   # 22nd and 24th succeeded
        self.assertEqual(len(result.failures), 1)
        self.assertEqual(result.failures[0].session_date, date(2026, 7, 23))
        self.assertIn("simulated fetch failure", result.failures[0].reason)
        self.assertIn("RuntimeError", result.failures[0].exception_text)

    def test_empty_day_produces_zero_trade_report_not_a_failure(self) -> None:
        result = run_multi_day_backtest([date(2026, 7, 22)], lambda d: _empty_day(d),
                                         STRIKE_GAP, NUM_STRIKES)
        self.assertEqual(len(result.failures), 0)
        self.assertEqual(result.daily_reports[0].trades, 0)
        self.assertEqual(result.daily_reports[0].win_rate_pct, 0.0)

    def test_reproducible_running_twice_gives_identical_results(self) -> None:
        days = [date(2026, 7, 22)]
        r1 = run_multi_day_backtest(days, lambda d: _one_trade_day(d), STRIKE_GAP, NUM_STRIKES)
        r2 = run_multi_day_backtest(days, lambda d: _one_trade_day(d), STRIKE_GAP, NUM_STRIKES)
        self.assertEqual(r1.daily_reports, r2.daily_reports)
        self.assertEqual([t.pnl for t in r1.all_trades], [t.pnl for t in r2.all_trades])

    def test_level_records_include_used_and_active_levels(self) -> None:
        result = run_multi_day_backtest([date(2026, 7, 22)], lambda d: _one_trade_day(d),
                                         STRIKE_GAP, NUM_STRIKES)
        states = {(strike, rec.final_state.value) for _d, rec in result.level_records
                  for strike in [rec.strike] if rec.strike == 24050}
        self.assertIn((24050, "USED"), states)   # the level that traded
        report = result.daily_reports[0]
        self.assertEqual(report.completed_mapped_levels, 1)
        self.assertGreater(report.unused_mapped_levels, 0)


class ExportToExcelTests(unittest.TestCase):
    def test_all_five_sheets_created_with_expected_headers(self) -> None:
        result = run_multi_day_backtest(
            [date(2026, 7, 22), date(2026, 7, 23)], lambda d: _one_trade_day(d),
            STRIKE_GAP, NUM_STRIKES,
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "backtest_summary.xlsx"
            export_to_excel(result, path)
            wb = load_workbook(path)
            self.assertEqual(
                wb.sheetnames,
                ["Daily Summary", "All Trades", "Mapped Level Statistics",
                 "Rejected Signals", "Validation Summary"],
            )
            self.assertEqual(wb["Daily Summary"].max_row, 3)   # header + 2 days
            self.assertEqual(wb["All Trades"].max_row, 3)       # header + 2 trades

    def test_failures_appear_in_validation_summary(self) -> None:
        def fetch_day(d):
            if d == date(2026, 7, 23):
                raise RuntimeError("boom")
            return _one_trade_day(d)

        result = run_multi_day_backtest(
            [date(2026, 7, 22), date(2026, 7, 23)], fetch_day, STRIKE_GAP, NUM_STRIKES,
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "backtest_summary.xlsx"
            export_to_excel(result, path)
            wb = load_workbook(path)
            values = [cell.value for row in wb["Validation Summary"].iter_rows() for cell in row]
            self.assertIn("2026-07-23", values)
            self.assertIn("boom", values)


if __name__ == "__main__":
    unittest.main()
