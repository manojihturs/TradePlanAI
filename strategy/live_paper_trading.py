"""Module 11: Live Paper Trading.

Single responsibility: run the EXISTING strategy engine (Modules 1-9,
unmodified) against a live stream of 5-minute candles, tracking a
paper trading capital ledger, dashboard state, and event log. This
module places NO real orders and has NO network dependency of its
own - live market data is delivered to it one closed candle at a time
by an external caller (a separate glue script owns the Upstox
connection), exactly like ``replay_engine.py`` is delivered historical
candles one at a time.

No entry, exit, mapping, level-state, or confirmation logic is added,
changed, or removed here. Trailing Stop Loss is reported as a field
for completeness, but Module 4 (Exit Signal) has no trailing-stop
behavior in the frozen v1.0 baseline - so the TSL field always equals
the Mapped Stop Loss, honestly reflecting "no trailing logic exists
yet" rather than inventing one.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Dict, List, Optional, Union

from openpyxl import Workbook

from strategy.entry_signal import Candle, EntrySignalDetector, TradeSide
from strategy.exit_signal import ExitLevels
from strategy.ladder_expansion import ensure_exit_levels
from strategy.level_capture import FirstCandleFetcher, LevelCapture
from strategy.level_state_manager import LevelState, LevelStateManager, level_key_from_entry
from strategy.paper_trading import PaperTradingEngine, Trade
from strategy.position_manager import Position, PositionManager
from strategy.premium_mapping import PremiumMapping

logger = logging.getLogger("live_paper_trading")


class LivePaperTradingError(Exception):
    """Raised for setup/usage errors in this module."""


# ---------------------------------------------------------------------------
# Capital tracking (paper account only - no real money, no real orders)
# ---------------------------------------------------------------------------

@dataclass
class CapitalTracker:
    """Tracks a paper trading account's capital and PnL.

    Attributes:
        initial_capital: The account's starting capital (never changes).
        lot_size: Contracts per trade, used to convert premium points to
            rupees. Not defined anywhere in the frozen Modules 1-10
            baseline (which works entirely in premium points) - this is
            an operational parameter Module 11 needs to report rupee
            capital figures, not a strategy rule. Defaults to NIFTY's
            existing lot size already used elsewhere in this project's
            legacy code (orb_common.QTY), for consistency.
    """

    initial_capital: float = 50_000.0
    lot_size: int = 65

    available_capital: float = field(init=False)
    used_capital: float = field(default=0.0, init=False)
    running_pnl: float = field(default=0.0, init=False)
    daily_pnl: float = field(default=0.0, init=False)
    total_pnl: float = field(default=0.0, init=False)
    _peak_total_pnl: float = field(default=0.0, init=False, repr=False)
    max_drawdown: float = field(default=0.0, init=False)

    def __post_init__(self) -> None:
        self.available_capital = self.initial_capital

    def reserve_for_entry(self, entry_premium: float) -> None:
        """Reserve capital when a paper position opens.

        Args:
            entry_premium: The entry premium (points) for the new position.

        Raises:
            LivePaperTradingError: if available capital is insufficient -
                reported, not silently allowed (still paper money, but
                the capital ledger must stay honest).
        """
        cost = entry_premium * self.lot_size
        if cost > self.available_capital:
            raise LivePaperTradingError(
                f"insufficient paper capital: need {cost:.2f}, have "
                f"{self.available_capital:.2f}"
            )
        self.available_capital -= cost
        self.used_capital += cost

    def release_on_exit(self, entry_premium: float, exit_premium: float) -> float:
        """Release reserved capital when a position closes and apply PnL.

        Args:
            entry_premium: The entry premium (points) that was reserved.
            exit_premium: The exit premium (points).

        Returns:
            The realized PnL in rupees for this trade.
        """
        cost = entry_premium * self.lot_size
        pnl_rupees = (exit_premium - entry_premium) * self.lot_size
        self.used_capital -= cost
        self.available_capital += cost + pnl_rupees
        self.running_pnl += pnl_rupees
        self.daily_pnl += pnl_rupees
        self.total_pnl += pnl_rupees
        self._peak_total_pnl = max(self._peak_total_pnl, self.total_pnl)
        self.max_drawdown = max(self.max_drawdown, self._peak_total_pnl - self.total_pnl)
        return pnl_rupees

    def reset_daily_pnl(self) -> None:
        """Reset the daily PnL counter - call once at the start of each day."""
        self.daily_pnl = 0.0


# ---------------------------------------------------------------------------
# Live trade record (per specification's Trade Management fields)
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class LiveTradeRecord:
    """One paper trade's full lifecycle record.

    Attributes:
        entry_time, entry_premium, target, stop_loss: As computed by
            Modules 3/4 at entry.
        trailing_stop_loss: Always equal to ``stop_loss`` - the frozen
            v1.0 baseline (Modules 1-10) has no trailing-stop logic,
            so this field honestly reports that rather than inventing
            trailing behavior (per instruction).
        exit_time, exit_premium, exit_reason: As computed by Module 4
            when the position closes; ``None`` while still open.
        running_pnl_rupees: PnL for this trade in rupees, using the
            configured lot size. ``None`` while still open.
    """

    strike: int
    side: TradeSide
    entry_time: datetime
    entry_premium: float
    target: float
    stop_loss: float
    trailing_stop_loss: float
    exit_time: Optional[datetime] = None
    exit_premium: Optional[float] = None
    exit_reason: Optional[str] = None
    running_pnl_rupees: Optional[float] = None


@dataclass(frozen=True)
class DashboardSnapshot:
    """A point-in-time view for continuous display, per specification."""

    current_time: datetime
    spot_price: Optional[float]
    top_strike: int
    bottom_strike: int
    current_open_trade: Optional[LiveTradeRecord]
    available_capital: float
    running_pnl: float
    todays_trades: int
    win_pct: float
    open_positions: int


@dataclass(frozen=True)
class RejectedSignalLog:
    timestamp: datetime
    strike: int
    side: TradeSide
    reason: str


# ---------------------------------------------------------------------------
# The live engine itself
# ---------------------------------------------------------------------------

class LivePaperTradingEngine:
    """Drives Modules 3, 5, 6, 8, 9 against a live stream of closed candles.

    This class performs NO network I/O - candles are pushed to it via
    ``on_candle_close`` by an external caller that owns the live data
    connection. No entry/exit/mapping/level-state logic is implemented
    here; every trading decision is delegated to the unmodified Modules
    1-9, in the same call sequence ``replay_engine.py`` and
    ``multi_day_backtest.py`` already use.
    """

    def __init__(
        self, session_date: date, capture: LevelCapture, mapping: PremiumMapping,
        strike_gap: int, fetch_first_candle: FirstCandleFetcher,
        capital: Optional[CapitalTracker] = None,
    ) -> None:
        self.session_date = session_date
        self.capture = capture
        self.mapping = mapping
        self.original_mapping = mapping
        self.strike_gap = strike_gap
        self.fetch_first_candle = fetch_first_candle

        self.detector = EntrySignalDetector(mapping)
        self.manager = PositionManager()
        self.ledger = PaperTradingEngine()
        self.level_states = LevelStateManager()
        self.level_states.initialize_day(mapping)
        self.capital = capital or CapitalTracker()

        self.trade_records: List[LiveTradeRecord] = []
        self.rejected_signals: List[RejectedSignalLog] = []
        self._pending_exit_levels: Optional[ExitLevels] = None
        self._current_live_record: Optional[LiveTradeRecord] = None

        logger.info("Opening Range Complete: ATM=%d Top=%.2f Bottom=%.2f",
                    capture.atm, capture.top_strike, capture.bottom_strike)
        logger.info("Top/Bottom Calculated: top_anchor=%d bottom_anchor=%d",
                    mapping.top_strike_rounded, mapping.bottom_strike_rounded)

    def on_candle_close(
        self, ts: datetime, ce_candles: Dict[int, Candle], pe_candles: Dict[int, Candle],
    ) -> None:
        """Process one newly-closed live 5-minute candle.

        Mirrors ``replay_engine.run_replay``'s per-candle loop exactly
        (same calls, same order) - see module docstring.
        """
        if self.manager.is_open:
            position = self.manager.current_position
            side_candles = ce_candles if position.entry.side is TradeSide.CE else pe_candles
            candle = side_candles.get(position.entry.strike)
            if candle is not None:
                self._pending_exit_levels = position.exit_levels
                closed = self.manager.process_candle(ts, candle)
                if closed is not None:
                    pnl_rupees = self.capital.release_on_exit(
                        (closed.entry.ce_level if closed.entry.side is TradeSide.CE
                         else closed.entry.pe_level),
                        closed.exit.exit_price,
                    )
                    self.ledger.record_trade(closed, self._pending_exit_levels)
                    self._pending_exit_levels = None
                    self.level_states.mark_used(level_key_from_entry(closed.entry))
                    if self._current_live_record is not None:
                        self.trade_records[-1] = LiveTradeRecord(
                            strike=self._current_live_record.strike,
                            side=self._current_live_record.side,
                            entry_time=self._current_live_record.entry_time,
                            entry_premium=self._current_live_record.entry_premium,
                            target=self._current_live_record.target,
                            stop_loss=self._current_live_record.stop_loss,
                            trailing_stop_loss=self._current_live_record.stop_loss,
                            exit_time=closed.exit.timestamp,
                            exit_premium=closed.exit.exit_price,
                            exit_reason=closed.exit.reason.value,
                            running_pnl_rupees=pnl_rupees,
                        )
                        self._current_live_record = None
                    logger.info(
                        "Trade Closed: %s strike=%d exit=%.2f reason=%s pnl_rupees=%.2f",
                        closed.entry.side.value, closed.entry.strike,
                        closed.exit.exit_price, closed.exit.reason.value, pnl_rupees,
                    )

        if self.manager.is_flat:
            signals = self.detector.process_candle(ts, ce_candles, pe_candles)
            chosen = None
            for signal in signals:
                key = level_key_from_entry(signal)
                if not self.level_states.can_open(key):
                    self.rejected_signals.append(RejectedSignalLog(
                        timestamp=ts, strike=signal.strike, side=signal.side,
                        reason="level already completed",
                    ))
                    logger.info("Signal Rejected: %s strike=%d reason=level already completed",
                                signal.side.value, signal.strike)
                    continue
                if chosen is None:
                    chosen = signal
                else:
                    self.rejected_signals.append(RejectedSignalLog(
                        timestamp=ts, strike=signal.strike, side=signal.side,
                        reason="already in position",
                    ))
                    logger.info("Signal Rejected: %s strike=%d reason=already in position",
                                signal.side.value, signal.strike)
            if chosen is not None:
                _levels, self.capture, self.mapping = ensure_exit_levels(
                    chosen, self.capture, self.mapping, self.strike_gap, self.fetch_first_candle,
                )
                position = self.manager.open(chosen, self.mapping)
                entry_premium = (chosen.ce_level if chosen.side is TradeSide.CE else chosen.pe_level)
                self.capital.reserve_for_entry(entry_premium)
                record = LiveTradeRecord(
                    strike=chosen.strike, side=chosen.side, entry_time=chosen.timestamp,
                    entry_premium=entry_premium, target=position.exit_levels.target,
                    stop_loss=position.exit_levels.stop_loss,
                    trailing_stop_loss=position.exit_levels.stop_loss,
                )
                self.trade_records.append(record)
                self._current_live_record = record
                logger.info(
                    "Trade Opened: %s strike=%d entry=%.2f target=%.2f stop=%.2f",
                    chosen.side.value, chosen.strike, entry_premium,
                    position.exit_levels.target, position.exit_levels.stop_loss,
                )
        else:
            self.detector.process_candle(ts, ce_candles, pe_candles)

    def dashboard_snapshot(self, current_time: datetime, spot_price: Optional[float]) -> DashboardSnapshot:
        """Build a point-in-time dashboard view - no side effects."""
        closed_trades = self.ledger.trades
        wins = sum(1 for t in closed_trades if t.pnl > 0)
        win_pct = (wins / len(closed_trades) * 100.0) if closed_trades else 0.0
        return DashboardSnapshot(
            current_time=current_time, spot_price=spot_price,
            top_strike=self.mapping.top_strike_rounded, bottom_strike=self.mapping.bottom_strike_rounded,
            current_open_trade=self._current_live_record if self.manager.is_open else None,
            available_capital=self.capital.available_capital, running_pnl=self.capital.running_pnl,
            todays_trades=len(closed_trades), win_pct=win_pct,
            open_positions=1 if self.manager.is_open else 0,
        )

    def end_of_day_summary(self) -> str:
        """A human-readable end-of-day summary, for logging."""
        trades = self.ledger.trades
        wins = sum(1 for t in trades if t.pnl > 0)
        losses = sum(1 for t in trades if t.pnl <= 0)
        return (
            f"End of Day Summary: {self.session_date} - {len(trades)} trades, "
            f"{wins}W/{losses}L, running_pnl_rupees={self.capital.running_pnl:.2f}, "
            f"total_pnl_rupees={self.capital.total_pnl:.2f}, "
            f"max_drawdown_rupees={self.capital.max_drawdown:.2f}, "
            f"available_capital={self.capital.available_capital:.2f}, "
            f"rejected_signals={len(self.rejected_signals)}"
        )


def export_end_of_day_excel(engine: LivePaperTradingEngine, path: Union[str, Path]) -> Path:
    """Export the 5-sheet end-of-day report, per specification.

    Sheets: Trade Log, Daily Summary, Level Usage, Rejected Signals,
    Capital Curve.
    """
    path = Path(path)
    wb = Workbook()

    trade_log = wb.active
    trade_log.title = "Trade Log"
    trade_log.append(["Strike", "Side", "Entry Time", "Entry Premium", "Target", "Stop Loss",
                       "Trailing Stop Loss", "Exit Time", "Exit Premium", "Exit Reason", "PnL (Rs)"])
    for r in engine.trade_records:
        trade_log.append([
            r.strike, r.side.value, r.entry_time.isoformat(), r.entry_premium, r.target,
            r.stop_loss, r.trailing_stop_loss,
            r.exit_time.isoformat() if r.exit_time else "", r.exit_premium or "",
            r.exit_reason or "OPEN", r.running_pnl_rupees if r.running_pnl_rupees is not None else "",
        ])

    trades = engine.ledger.trades
    wins = sum(1 for t in trades if t.pnl > 0)
    losses = sum(1 for t in trades if t.pnl <= 0)
    daily = wb.create_sheet("Daily Summary")
    daily.append(["Metric", "Value"])
    daily.append(["Date", engine.session_date.isoformat()])
    daily.append(["Trades", len(trades)])
    daily.append(["Wins", wins])
    daily.append(["Losses", losses])
    daily.append(["Win %", round((wins / len(trades) * 100.0) if trades else 0.0, 2)])
    daily.append(["Running PnL (Rs)", engine.capital.running_pnl])
    daily.append(["Total PnL (Rs)", engine.capital.total_pnl])
    daily.append(["Max Drawdown (Rs)", engine.capital.max_drawdown])
    daily.append(["Available Capital (Rs)", engine.capital.available_capital])
    daily.append(["Used Capital (Rs)", engine.capital.used_capital])

    level_usage = wb.create_sheet("Level Usage")
    level_usage.append(["Strike", "Anchor", "Side", "Final State"])
    all_keys = set()
    for strike in engine.original_mapping.top_ce_ladder:
        all_keys.add((strike, "TOP", "CE"))
    for strike in engine.original_mapping.top_pe_ladder:
        all_keys.add((strike, "TOP", "PE"))
    for strike in engine.original_mapping.bottom_ce_ladder:
        all_keys.add((strike, "BOTTOM", "CE"))
    for strike in engine.original_mapping.bottom_pe_ladder:
        all_keys.add((strike, "BOTTOM", "PE"))
    from strategy.level_state_manager import LevelKey, MappingAnchor
    for strike, anchor_s, side_s in sorted(all_keys):
        anchor = MappingAnchor.TOP if anchor_s == "TOP" else MappingAnchor.BOTTOM
        side = TradeSide.CE if side_s == "CE" else TradeSide.PE
        state = engine.level_states.get_state(LevelKey(strike, anchor, side))
        level_usage.append([strike, anchor_s, side_s, state.value])

    rejected = wb.create_sheet("Rejected Signals")
    rejected.append(["Time", "Strike", "Side", "Reason"])
    for r in engine.rejected_signals:
        rejected.append([r.timestamp.isoformat(), r.strike, r.side.value, r.reason])

    capital_curve = wb.create_sheet("Capital Curve")
    capital_curve.append(["Trade #", "Exit Time", "Cumulative PnL (Rs)", "Available Capital (Rs)"])
    running = 0.0
    available = engine.capital.initial_capital
    for i, r in enumerate(engine.trade_records, start=1):
        if r.running_pnl_rupees is not None:
            running += r.running_pnl_rupees
            available += r.running_pnl_rupees
            capital_curve.append([i, r.exit_time.isoformat() if r.exit_time else "", running, available])

    wb.save(path)
    logger.info("Exported end-of-day report to %s", path)
    return path
