"""Post-market workflow: data loading.

Single responsibility: locate the latest completed trading session
produced by Module 11 (``strategy/live_paper_trading.py`` via
``run_live_paper_trading.py``) and load its Excel export and log file
into plain in-memory records. Performs NO analysis - that is left to
the other ``postmarket`` modules. This is the only module in the
package that touches the filesystem for reading session data.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook

logger = logging.getLogger("postmarket.data_sources")

_REPORT_FILENAME_RE = re.compile(r"live_paper_trading_(\d{4}-\d{2}-\d{2})\.xlsx$")
_ATM_LOG_RE = re.compile(r"Opening Range Complete: ATM=(\d+)")


class DataSourceError(Exception):
    """Raised when the latest session's data cannot be located or read."""


@dataclass(frozen=True)
class TradeRecord:
    strike: int
    side: str
    entry_time: Optional[datetime]
    entry_premium: float
    target: float
    stop_loss: float
    trailing_stop_loss: float
    exit_time: Optional[datetime]
    exit_premium: Optional[float]
    exit_reason: str
    pnl_rupees: Optional[float]


@dataclass(frozen=True)
class RejectedSignalRecord:
    timestamp: Optional[datetime]
    strike: int
    side: str
    reason: str


@dataclass(frozen=True)
class LevelUsageRecord:
    strike: int
    anchor: str
    side: str
    final_state: str


@dataclass
class SessionData:
    """Everything the post-market workflow needs for one trading day."""

    session_date: date
    xlsx_path: Path
    log_path: Optional[Path]
    atm: Optional[int]
    daily_summary: Dict[str, object] = field(default_factory=dict)
    trades: List[TradeRecord] = field(default_factory=list)
    rejected_signals: List[RejectedSignalRecord] = field(default_factory=list)
    level_usage: List[LevelUsageRecord] = field(default_factory=list)
    log_lines: List[str] = field(default_factory=list)


def find_latest_session_report(reports_dir: Path) -> Path:
    """Find the most recent ``live_paper_trading_<date>.xlsx`` report.

    Args:
        reports_dir: Directory to search (non-recursive).

    Returns:
        Path to the report with the latest ``session_date`` in its name.

    Raises:
        DataSourceError: if no matching report file exists.
    """
    candidates = []
    for p in reports_dir.glob("live_paper_trading_*.xlsx"):
        m = _REPORT_FILENAME_RE.search(p.name)
        if m:
            candidates.append((date.fromisoformat(m.group(1)), p))
    if not candidates:
        raise DataSourceError(
            f"no live_paper_trading_<date>.xlsx report found in {reports_dir}"
        )
    candidates.sort(key=lambda t: t[0])
    latest_date, latest_path = candidates[-1]
    logger.info("Latest completed session found: %s (%s)", latest_date, latest_path)
    return latest_path


def _parse_dt(value: object) -> Optional[datetime]:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def load_session_report(xlsx_path: Path, log_path: Optional[Path] = None) -> SessionData:
    """Load a Module 11 end-of-day Excel export into plain records.

    Args:
        xlsx_path: Path to the ``live_paper_trading_<date>.xlsx`` export
            (5 sheets: Trade Log, Daily Summary, Level Usage, Rejected
            Signals, Capital Curve - see
            ``strategy.live_paper_trading.export_end_of_day_excel``).
        log_path: Optional path to the matching run log, used for ATM
            extraction and API/data health scanning.

    Returns:
        A populated ``SessionData``.

    Raises:
        DataSourceError: if the workbook is missing an expected sheet.
    """
    m = _REPORT_FILENAME_RE.search(xlsx_path.name)
    if not m:
        raise DataSourceError(f"cannot parse session date from filename: {xlsx_path.name}")
    session_date = date.fromisoformat(m.group(1))

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    for sheet_name in ("Trade Log", "Daily Summary", "Level Usage", "Rejected Signals"):
        if sheet_name not in wb.sheetnames:
            raise DataSourceError(f"{xlsx_path} is missing expected sheet '{sheet_name}'")

    trades: List[TradeRecord] = []
    rows = list(wb["Trade Log"].iter_rows(min_row=2, values_only=True))
    for row in rows:
        if row[0] is None:
            continue
        (strike, side, entry_time, entry_premium, target, stop_loss, tsl,
         exit_time, exit_premium, exit_reason, pnl) = row
        trades.append(TradeRecord(
            strike=int(strike), side=str(side),
            entry_time=_parse_dt(entry_time), entry_premium=float(entry_premium),
            target=float(target), stop_loss=float(stop_loss),
            trailing_stop_loss=float(tsl),
            exit_time=_parse_dt(exit_time),
            exit_premium=float(exit_premium) if exit_premium not in (None, "") else None,
            exit_reason=str(exit_reason) if exit_reason else "OPEN",
            pnl_rupees=float(pnl) if pnl not in (None, "") else None,
        ))

    daily_summary: Dict[str, object] = {}
    for row in wb["Daily Summary"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        daily_summary[str(row[0])] = row[1]

    level_usage: List[LevelUsageRecord] = []
    for row in wb["Level Usage"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        strike, anchor, side, final_state = row
        level_usage.append(LevelUsageRecord(
            strike=int(strike), anchor=str(anchor), side=str(side), final_state=str(final_state),
        ))

    rejected_signals: List[RejectedSignalRecord] = []
    for row in wb["Rejected Signals"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        ts, strike, side, reason = row
        rejected_signals.append(RejectedSignalRecord(
            timestamp=_parse_dt(ts), strike=int(strike), side=str(side), reason=str(reason),
        ))

    wb.close()

    log_lines: List[str] = []
    atm: Optional[int] = None
    if log_path is not None and log_path.exists():
        log_lines = log_path.read_text(encoding="utf-8", errors="replace").splitlines()
        for line in log_lines:
            m_atm = _ATM_LOG_RE.search(line)
            if m_atm:
                atm = int(m_atm.group(1))
                break

    return SessionData(
        session_date=session_date, xlsx_path=xlsx_path, log_path=log_path, atm=atm,
        daily_summary=daily_summary, trades=trades, rejected_signals=rejected_signals,
        level_usage=level_usage, log_lines=log_lines,
    )
