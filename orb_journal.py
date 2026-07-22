# orb_journal.py - per-day trade journal in a single Excel workbook.
# Each trading day gets its own sheet (named YYYY-MM-DD) with one row per
# trade, including a human-readable note on why the trade was taken and
# why it made or lost money.

import os
from openpyxl import Workbook, load_workbook

_HERE = os.path.dirname(os.path.abspath(__file__))
JOURNAL_PATH = os.path.join(_HERE, "trade_journal.xlsx")

HEADERS = ["entry_ts", "exit_ts", "side", "atm_strike", "strike", "qty", "entry_price",
           "exit_price", "pnl_points", "pnl_rupees", "lines_crossed",
           "exit_reason", "trigger_strike", "trigger_strike_value",
           "why_entered", "why_pnl"]

def _open_workbook():
    if os.path.exists(JOURNAL_PATH):
        return load_workbook(JOURNAL_PATH)
    wb = Workbook()
    wb.remove(wb.active)   # drop the default blank sheet
    return wb

def append_trade(session_date, row):
    """row: dict with keys matching HEADERS (missing keys become blank)."""
    sheet_name = str(session_date)[:31]   # Excel sheet name length limit
    wb = _open_workbook()
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
        ws.append(HEADERS)
    else:
        ws = wb[sheet_name]
    ws.append([row.get(h, "") for h in HEADERS])
    wb.save(JOURNAL_PATH)
